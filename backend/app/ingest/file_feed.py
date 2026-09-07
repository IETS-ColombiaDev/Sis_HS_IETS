"""Adaptador generico de descargas estructuradas (CSV / XLSX).

TGA es el primer usuario. El mapeo de columnas se declara en
`connector_config.column_map` para no acoplar el parseo a una agencia.
"""
from __future__ import annotations

import csv
import io

from .base import (
    CanonicalRecord,
    Connector,
    ConnectorError,
    ConnectorResult,
    clean_text,
    parse_compact_date,
    register,
    request_bytes,
    schema_signature,
)


class FileFeedConnector(Connector):
    code = "file_feed"
    label = "Descarga estructurada (CSV / XLSX)"
    description = "Baja un archivo oficial y lo traduce con el mapa de columnas de la fuente."
    min_interval = 1.0
    adapter_version = "1"

    def fetch(self, *, config: dict, url: str = "") -> ConnectorResult:
        config = config or {}
        if config.get("records"):
            rows = [r for r in config["records"] if isinstance(r, dict)]
            records = [self._to_record(row, config) for row in rows]
            records = [r for r in records if r]
            return ConnectorResult(
                records=records,
                message=f"{len(records)} filas cargadas desde lote local.",
                adapter_version=self.adapter_version,
            )

        endpoint = clean_text(config.get("download_url") or config.get("file_url") or url)
        if not endpoint:
            raise ConnectorError(
                "file_feed necesita download_url o un lote en connector_config.records. "
                "Para TGA, descargue el CSV/XLSX del ARTG y carguelo como lote."
            )

        payload, content_type = request_bytes(
            endpoint,
            connector_code=self.code,
            min_interval=self.min_interval,
        )
        rows = _parse_tabular(payload, content_type, config)
        records = [self._to_record(row, config) for row in rows]
        records = [r for r in records if r]
        return ConnectorResult(
            records=records,
            message=f"{len(records)} filas del archivo estructurado.",
            schema_signature=schema_signature(rows[:3]),
            adapter_version=self.adapter_version,
            endpoint=endpoint,
        )

    def _to_record(self, row: dict, config: dict) -> CanonicalRecord | None:
        mapped = _apply_map(row, config.get("column_map") or {})
        external_id = clean_text(mapped.get("external_id") or row.get("id") or "")
        title = clean_text(mapped.get("title") or mapped.get("commercial_name") or "", 590)
        if not external_id and not title:
            return None
        approval = parse_compact_date(mapped.get("approval_date") or mapped.get("published_date"))
        return CanonicalRecord(
            external_id=external_id or title[:80],
            title=title or external_id,
            url=clean_text(mapped.get("url") or "", 1000),
            summary=clean_text(mapped.get("summary") or mapped.get("indication") or title, 1500),
            commercial_name=clean_text(mapped.get("commercial_name") or title, 400),
            inn_name=clean_text(mapped.get("inn_name"), 400),
            manufacturer=clean_text(mapped.get("manufacturer"), 300),
            indication=clean_text(mapped.get("indication"), 2000),
            technology_type=clean_text(mapped.get("technology_type") or "medicamento"),
            development_phase=clean_text(mapped.get("development_phase") or "Registro"),
            horizon="inminente" if approval else "transicional",
            fda_approval_date=None,
            ema_approval_date=None,
            regulatory_status=clean_text(mapped.get("regulatory_status") or "Registro TGA / file_feed"),
            published_date=clean_text(mapped.get("published_date")),
            raw=row,
        )


def _apply_map(row: dict, column_map: dict) -> dict:
    out: dict[str, str] = {}
    lowered = {str(k).strip().lower(): v for k, v in row.items()}
    for target, candidates in column_map.items():
        names = candidates if isinstance(candidates, list) else [candidates]
        for name in names:
            value = row.get(name)
            if value in (None, ""):
                value = lowered.get(str(name).strip().lower())
            if value not in (None, ""):
                out[target] = value
                break
    for key, value in row.items():
        if key not in out and value not in (None, ""):
            out.setdefault(str(key), value)
    return out


def _parse_tabular(payload: bytes, content_type: str, config: dict) -> list[dict]:
    fmt = (config.get("format") or "").lower()
    name_hint = (config.get("download_url") or config.get("file_url") or "").lower()
    if fmt == "xlsx" or "spreadsheet" in content_type or name_hint.endswith(".xlsx"):
        return _parse_xlsx(payload)
    text = payload.decode("utf-8-sig", errors="replace")
    sample = text[:2048]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    reader = csv.DictReader(io.StringIO(text), dialect=dialect)
    return [dict(row) for row in reader if any((v or "").strip() for v in row.values())]


def _parse_xlsx(payload: bytes) -> list[dict]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ConnectorError(
            "Para leer XLSX instale openpyxl (ya esta en requirements.txt)."
        ) from exc
    book = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
    sheet = book.active
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [str(h or f"col_{i}").strip() for i, h in enumerate(rows[0])]
    out = []
    for raw in rows[1:]:
        item = {headers[i]: ("" if raw[i] is None else raw[i]) for i in range(min(len(headers), len(raw)))}
        if any(str(v).strip() for v in item.values()):
            out.append(item)
    return out


register(FileFeedConnector())
