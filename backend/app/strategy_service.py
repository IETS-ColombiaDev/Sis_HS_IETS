"""Datamart estrategico, ficha publica, boletines y alertas (fase 6)."""
from __future__ import annotations

import re
from datetime import date, datetime, timedelta, timezone
from decimal import ROUND_HALF_UP, Decimal

from sqlalchemy import func, or_
from sqlalchemy.orm import Session

from . import audit, ttm
from .methodology import TECHNOLOGY_STATUS_LABELS, get_param
from .models import (
    AlertEvent,
    AlertSubscription,
    Bulletin,
    Cluster,
    Cycle,
    CycleDatamart,
    CycleTechnology,
    EvaluationDoc,
    TechType,
    Technology,
    TimeToMarketSnapshot,
    User,
)

PUBLIC_HIDDEN_FIELDS = {
    "budget_year_1",
    "budget_year_2",
    "budget_year_3",
    "comparators_sgsss",
    "early_dialogue_notes",
    "confidential_note",
}

BULLETIN_STATUSES = ("borrador", "pendiente_aprobacion", "publicado")


class StrategyRuleError(ValueError):
    pass


def _utcnow() -> datetime:
    return datetime.now(timezone.utc)


def _tech_name(tech: Technology | None) -> str:
    if tech is None:
        return "Tecnología"
    return (tech.commercial_name or tech.inn_name or f"Tecnología {tech.id}").strip()


PHASE_LABELS = {
    "autorizada": "Autorizada / en mercado",
    "fase_iv": "Fase IV",
    "fase_iii": "Fase III",
    "fase_ii": "Fase II",
    "fase_i": "Fase I",
    "sin_fase": "Sin fase declarada",
}

# Orden de lectura del tablero: de la mas lejana a la mas cercana al mercado.
PHASE_ORDER = ("fase_i", "fase_ii", "fase_iii", "fase_iv", "autorizada", "sin_fase")

# Embudo de conversion del ciclo (RF17). Cada etapa contiene a la siguiente: una
# tecnologia publicada tambien cuenta como evaluada, priorizada, filtrada y
# capturada. El estado que manda es el de la instancia del ciclo
# (`CycleTechnology.status`) y el expediente del mismo ciclo, nunca el estado
# global de la tecnologia, que puede venir de otro ciclo.
FUNNEL_STAGES = ("captured", "filtered", "prioritized", "evaluated", "published")
FUNNEL_LABELS = {
    "captured": "Capturadas",
    "filtered": "Filtradas",
    "prioritized": "Priorizadas",
    "evaluated": "Evaluadas",
    "published": "Publicadas",
}
FILTERED_STATUSES = frozenset(
    {
        "filtrada_apta_priorizacion",
        "priorizada",
        "bajo_vigilancia",
        "no_priorizada",
        "en_evaluacion",
        "publicada",
    }
)
PRIORITIZED_STATUSES = frozenset({"priorizada", "en_evaluacion", "publicada"})
EVALUATED_STATUSES = frozenset({"en_evaluacion", "publicada"})

# Campos de la fila que solo viajan con `analytics:restricted`.
RESTRICTED_ROW_KEYS = ("year1", "year2", "year3", "budget_unparsed", "comparators")

# Version del formato del datamart. La 2 guarda filas de hechos por tecnologia,
# de modo que los filtros de un ciclo cerrado se resuelven sobre el datamart y no
# sobre las tablas transaccionales.
DATAMART_SCHEMA = 2

# Colombia no aplica horario de verano: UTC-5 fijo.
COLOMBIA_TZ = timezone(timedelta(hours=-5))

_ACCENTS = str.maketrans("áéíóúüñÁÉÍÓÚÜÑ", "aeiouunAEIOUUN")


def _plain(text: str) -> str:
    return (text or "").translate(_ACCENTS).lower()


# --------------------------------------------------------------------------- #
#  Montos presupuestales
# --------------------------------------------------------------------------- #
_MONEY_RE = re.compile(
    r"(?P<cur>\$|cop|usd|us\$)?\s*"
    r"(?P<num>\d(?:[\d.,]*\d)?)"
    r"(?:\s*(?P<mult>mil\s+millones|millones|millon|billones|billon|mmm|mm|mil)\b)?",
    re.IGNORECASE,
)
_MULTIPLIERS = {
    "mil millones": 1e9,
    "millones": 1e6,
    "millon": 1e6,
    "billones": 1e12,
    "billon": 1e12,
    "mmm": 1e9,
    "mm": 1e6,
    "mil": 1e3,
}
# Por debajo de este valor, un numero sin moneda ni multiplicador no se toma
# como monto: suele ser un anio, un conteo de pacientes o un porcentaje.
MONEY_MIN_BARE = 1_000_000


def _parse_number(token: str) -> float | None:
    tok = token.strip()
    has_dot, has_comma = "." in tok, "," in tok
    try:
        if has_dot and has_comma:
            decimal = "." if tok.rfind(".") > tok.rfind(",") else ","
            thousands = "," if decimal == "." else "."
            return float(tok.replace(thousands, "").replace(decimal, "."))
        sep = "." if has_dot else ("," if has_comma else "")
        if not sep:
            return float(tok)
        parts = tok.split(sep)
        if len(parts) > 2:
            return float("".join(parts))
        head, tail = parts
        if len(tail) == 3 and head not in ("", "0"):
            return float(head + tail)
        return float(f"{head or '0'}.{tail}")
    except ValueError:
        return None


def parse_money(value) -> float | None:
    """Primer monto reconocible del texto, en la unidad escrita (COP por convencion).

    Entiende separadores colombianos e ingleses ("3.200.000.000", "3,200,000",
    "3,5 mil millones", "1.200 millones", "$ 45 MM"). Un numero pelado menor de
    un millon no se toma como monto. Devuelve None si no hay monto.
    """
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = _plain(str(value or ""))
    for match in _MONEY_RE.finditer(text):
        number = _parse_number(match.group("num"))
        if number is None:
            continue
        mult_key = re.sub(r"\s+", " ", (match.group("mult") or "").lower())
        amount = number * _MULTIPLIERS.get(mult_key, 1.0)
        if match.group("cur") or mult_key or amount >= MONEY_MIN_BARE:
            return float(amount)
    return None


def _parse_money(value) -> float:
    """Compatibilidad: 0.0 cuando no hay monto."""
    return parse_money(value) or 0.0


# --------------------------------------------------------------------------- #
#  Fase clinica
# --------------------------------------------------------------------------- #
_ROMAN = {"i": 1, "ii": 2, "iii": 3, "iv": 4, "1": 1, "2": 2, "3": 3, "4": 4}
_PHASE_BY_NUMBER = {1: "fase_i", 2: "fase_ii", 3: "fase_iii", 4: "fase_iv"}


def _phase_bucket(text: str) -> str:
    """Agrupa la fase declarada. Gana la fase mas avanzada que se mencione."""
    low = _plain(text)
    if not low.strip():
        return "sin_fase"
    negated = re.search(r"\b(no|not|sin|non)\s+(autoriz|aprob|approv|authori)", low)
    if not negated and re.search(r"autoriz|aprobad|approved|authori[sz]ed", low):
        return "autorizada"
    norm = low.replace("phase", "fase")
    norm = re.sub(r"fase\s*[_-]?\s*(\d)", r"fase \1", norm)
    found: set[int] = set()
    for match in re.finditer(r"fase\s+((?:iv|iii|ii|i|[1-4])\b(?:\s*[/|,-]\s*(?:fase\s+)?(?:iv|iii|ii|i|[1-4])\b)*)", norm):
        for token in re.findall(r"iv|iii|ii|i|[1-4]", match.group(1)):
            found.add(_ROMAN[token])
    if not found:
        if re.search(r"\biv\b", low):
            found.add(4)
        elif re.search(r"\biii\b", low):
            found.add(3)
        elif re.search(r"\bii\b", low):
            found.add(2)
    if not found:
        return "sin_fase"
    return _PHASE_BY_NUMBER[max(found)]


def _as_date(value) -> date | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    try:
        return date.fromisoformat(str(value)[:10])
    except ValueError:
        return None


def _local_date(value) -> date | None:
    """Fecha de captura en hora de Colombia (la base guarda UTC)."""
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        stamp = value if value.tzinfo else value.replace(tzinfo=timezone.utc)
        return stamp.astimezone(COLOMBIA_TZ).date()
    return _as_date(value)


def _pct(part: int, whole: int) -> int:
    """Porcentaje entero con redondeo comercial (0,5 sube)."""
    if not whole:
        return 0
    return int((Decimal(100 * int(part)) / Decimal(int(whole))).quantize(Decimal("1"), rounding=ROUND_HALF_UP))


def _conversion(funnel: dict) -> dict:
    captured = int(funnel.get("captured") or 0)
    return {
        "filter_rate": _pct(int(funnel.get("filtered") or 0), captured),
        "priority_rate": _pct(int(funnel.get("prioritized") or 0), captured),
        "eval_rate": _pct(int(funnel.get("evaluated") or 0), captured),
        "publish_rate": _pct(int(funnel.get("published") or 0), captured),
    }


def funnel_stage(entry_status: str, doc: EvaluationDoc | None) -> str:
    """Etapa mas avanzada del embudo que alcanzo la tecnologia en este ciclo."""
    if entry_status not in FILTERED_STATUSES:
        return "captured"
    if entry_status not in PRIORITIZED_STATUSES:
        return "filtered"
    if entry_status not in EVALUATED_STATUSES and doc is None:
        return "prioritized"
    if entry_status == "publicada" or (doc is not None and doc.status == "publicado"):
        return "published"
    return "evaluated"


def _looks_phase3(text: str) -> bool:
    low = (text or "").lower()
    return "iii" in low or "phase 3" in low or "fase 3" in low or "phase3" in low


def _mentions_colombia(db: Session, tech: Technology) -> bool:
    raw = get_param(db, "alerts.phase3_country_tokens", "colombia,colombiano,colombiana")
    tokens = [t.strip().lower() for t in str(raw or "").split(",") if t.strip()]
    haystack = " ".join(
        [
            tech.indication or "",
            tech.summary or "",
            tech.regulatory_status or "",
            tech.development_phase or "",
        ]
    ).lower()
    return any(tok in haystack for tok in tokens)


def compute_ttm_row(db: Session, tech: Technology) -> dict:
    result = ttm.compute_for(db, tech)
    result["technology_id"] = tech.id
    result["name"] = _tech_name(tech)
    return result


def snapshot_ttm(db: Session, cycle: Cycle) -> int:
    entries = db.query(CycleTechnology).filter(CycleTechnology.cycle_id == cycle.id).all()
    params = ttm.load_params(db)
    tech_ids = [e.technology_id for e in entries]
    techs = (
        {t.id: t for t in db.query(Technology).filter(Technology.id.in_(tech_ids)).all()}
        if tech_ids
        else {}
    )
    existing = {
        row.technology_id: row
        for row in db.query(TimeToMarketSnapshot).filter(TimeToMarketSnapshot.cycle_id == cycle.id).all()
    }
    written = 0
    for entry in entries:
        tech = techs.get(entry.technology_id)
        if tech is None:
            continue
        calc = ttm.compute_with(params, tech)
        row = existing.get(tech.id)
        if row is None:
            row = TimeToMarketSnapshot(cycle_id=cycle.id, technology_id=tech.id)
            db.add(row)
            existing[tech.id] = row
        row.months = calc["months"]
        row.band = calc["band"]
        row.basis = calc["basis"]
        row.computed_at = _utcnow()
        written += 1
    return written


# --------------------------------------------------------------------------- #
#  Filas de hechos del tablero
# --------------------------------------------------------------------------- #
def collect_rows(db: Session, cycle: Cycle, *, as_of: date | None = None) -> list[dict]:
    """Una fila por tecnologia del ciclo con todas las dimensiones del tablero.

    Es la unica fuente de los agregados, tanto en vivo como en el datamart. Para
    un ciclo cerrado el time-to-market se evalua en la fecha de su version
    (snapshot del cierre), no en la fecha de hoy.
    """
    entries = (
        db.query(CycleTechnology)
        .filter(CycleTechnology.cycle_id == cycle.id)
        .order_by(CycleTechnology.id)
        .all()
    )
    tech_ids = [e.technology_id for e in entries]
    techs = (
        {t.id: t for t in db.query(Technology).filter(Technology.id.in_(tech_ids)).all()}
        if tech_ids
        else {}
    )
    docs = {
        d.technology_id: d
        for d in db.query(EvaluationDoc).filter(EvaluationDoc.cycle_id == cycle.id).all()
    }
    clusters = {c.id: c.name for c in db.query(Cluster).all()}
    types = {t.id: t.name for t in db.query(TechType).all()}
    params = ttm.load_params(db)

    versioned: dict[int, date] = {}
    fallback_as_of = as_of
    if cycle.status == "cerrado_consolidado" and as_of is None:
        for tech_id, computed_at in (
            db.query(TimeToMarketSnapshot.technology_id, TimeToMarketSnapshot.computed_at)
            .filter(TimeToMarketSnapshot.cycle_id == cycle.id)
            .all()
        ):
            if computed_at is not None:
                versioned[tech_id] = _local_date(computed_at)
        if cycle.closed_at is not None:
            fallback_as_of = _local_date(cycle.closed_at)

    rows: list[dict] = []
    for entry in entries:
        tech = techs.get(entry.technology_id)
        if tech is None:
            continue
        doc = docs.get(tech.id)
        calc = ttm.compute_with(params, tech, as_of=versioned.get(tech.id) or fallback_as_of)
        phase_code = _phase_bucket(tech.development_phase)
        captured_on = _local_date(tech.captured_at)
        body = (doc.body or {}) if doc else {}
        budget = {}
        unparsed = False
        for year in (1, 2, 3):
            raw_text = body.get(f"budget_year_{year}")
            amount = parse_money(raw_text)
            if amount is None and str(raw_text or "").strip():
                unparsed = True
            budget[f"year{year}"] = float(amount or 0.0)
        rows.append(
            {
                "technology_id": tech.id,
                "name": _tech_name(tech),
                "cluster_id": tech.cluster_id,
                "cluster": clusters.get(tech.cluster_id, "Sin clúster"),
                "tech_type_id": tech.tech_type_id,
                "tech_type": types.get(tech.tech_type_id, "Sin tipología"),
                "status": entry.status,
                "status_label": TECHNOLOGY_STATUS_LABELS.get(entry.status, entry.status),
                "stage": funnel_stage(entry.status, doc),
                "points": int(entry.priority_points or 0),
                "points_rated": entry.priority_points is not None,
                "priority_pct": float(entry.priority_pct) if entry.priority_pct is not None else None,
                "phase": tech.development_phase or "",
                "phase_bucket": phase_code,
                "phase_label": PHASE_LABELS.get(phase_code, phase_code),
                "months": calc["months"],
                "months_raw": calc["months_raw"],
                "band": calc["band"],
                "band_label": calc["band_label"],
                "basis": calc["basis"],
                "basis_label": calc["basis_label"],
                "reference_date": calc["reference_date"],
                "approved": calc["approved"],
                "overdue": calc["overdue"],
                "captured_on": captured_on.isoformat() if captured_on else None,
                "doc_status": doc.status if doc else "",
                "product_level": doc.product_level if doc else "",
                "year1": budget["year1"],
                "year2": budget["year2"],
                "year3": budget["year3"],
                "budget_unparsed": unparsed,
                "comparators": str(body.get("comparators_sgsss") or "").strip(),
            }
        )
    return rows


def normalize_filters(
    *,
    cluster_id: int | None = None,
    tech_type_id: int | None = None,
    band: str = "",
    status: str = "",
    phase: str = "",
    date_from: date | None = None,
    date_to: date | None = None,
    priority_min: int | None = None,
    stage: str = "",
) -> dict:
    """Filtros efectivos. `cluster_id=0` o `tech_type_id=0` significan "sin asignar"."""
    return {
        "cluster_id": cluster_id,
        "tech_type_id": tech_type_id,
        "band": (band or "").strip(),
        "status": (status or "").strip(),
        "phase": (phase or "").strip(),
        "date_from": _as_date(date_from),
        "date_to": _as_date(date_to),
        "priority_min": priority_min,
        "stage": (stage or "").strip(),
    }


def _active(filters: dict) -> dict:
    return {k: v for k, v in filters.items() if v not in (None, "")}


def row_matches(row: dict, filters: dict) -> bool:
    cluster_id = filters.get("cluster_id")
    if cluster_id is not None:
        if cluster_id == 0:
            if row.get("cluster_id") is not None:
                return False
        elif row.get("cluster_id") != cluster_id:
            return False
    tech_type_id = filters.get("tech_type_id")
    if tech_type_id is not None:
        if tech_type_id == 0:
            if row.get("tech_type_id") is not None:
                return False
        elif row.get("tech_type_id") != tech_type_id:
            return False
    if filters.get("status") and row.get("status") != filters["status"]:
        return False
    if filters.get("band") and row.get("band") != filters["band"]:
        return False
    if filters.get("phase") and row.get("phase_bucket") != filters["phase"]:
        return False
    if filters.get("stage"):
        if FUNNEL_STAGES.index(row.get("stage") or "captured") < FUNNEL_STAGES.index(filters["stage"]):
            return False
    priority_min = filters.get("priority_min")
    if priority_min is not None:
        if not row.get("points_rated") or int(row.get("points") or 0) < int(priority_min):
            return False
    date_from, date_to = filters.get("date_from"), filters.get("date_to")
    if date_from or date_to:
        captured = _as_date(row.get("captured_on"))
        if captured is None:
            return False
        if date_from and captured < date_from:
            return False
        if date_to and captured > date_to:
            return False
    return True


def filter_rows(rows: list[dict], filters: dict) -> list[dict]:
    active = _active(filters)
    if not active:
        return list(rows)
    return [row for row in rows if row_matches(row, active)]


def _public_row(row: dict) -> dict:
    return {k: v for k, v in row.items() if k not in RESTRICTED_ROW_KEYS}


def aggregate(rows: list[dict]) -> dict:
    """Agregados del tablero a partir de las filas ya filtradas."""
    reached = {stage: 0 for stage in FUNNEL_STAGES}
    for row in rows:
        idx = FUNNEL_STAGES.index(row.get("stage") or "captured")
        for stage in FUNNEL_STAGES[: idx + 1]:
            reached[stage] += 1

    clusters: dict[tuple, int] = {}
    types: dict[tuple, int] = {}
    bands = {band: 0 for band in ttm.TTM_BANDS}
    phases = {phase: 0 for phase in PHASE_ORDER}
    for row in rows:
        ckey = (row.get("cluster_id"), row.get("cluster") or "Sin clúster")
        clusters[ckey] = clusters.get(ckey, 0) + 1
        tkey = (row.get("tech_type_id"), row.get("tech_type") or "Sin tipología")
        types[tkey] = types.get(tkey, 0) + 1
        bands[row.get("band") or "desconocido"] = bands.get(row.get("band") or "desconocido", 0) + 1
        phases[row.get("phase_bucket") or "sin_fase"] = phases.get(row.get("phase_bucket") or "sin_fase", 0) + 1

    def _ranked(counter: dict, id_key: str) -> list[dict]:
        items = [{"label": label, "value": n, id_key: ident} for (ident, label), n in counter.items()]
        return sorted(items, key=lambda i: (-i["value"], i["label"]))

    with_months = [r for r in rows if r.get("months") is not None]
    return {
        "funnel": reached,
        "conversion": _conversion(reached),
        "by_cluster": _ranked(clusters, "cluster_id"),
        "by_type": _ranked(types, "tech_type_id"),
        "by_band": [
            {"label": ttm.TTM_BAND_LABELS.get(code, code), "code": code, "value": bands[code]}
            for code in ttm.TTM_BANDS
        ],
        "by_phase": [
            {"label": PHASE_LABELS.get(code, code), "code": code, "value": phases[code]}
            for code in PHASE_ORDER
        ],
        "ttm_summary": {
            "total": len(rows),
            "with_estimate": len(with_months),
            "without_estimate": len(rows) - len(with_months),
            "already_approved": sum(1 for r in with_months if r.get("approved")),
            "overdue": sum(1 for r in with_months if r.get("overdue")),
        },
    }


def _budget(rows: list[dict]) -> dict:
    heat: dict[tuple, dict] = {}
    for row in rows:
        key = (row.get("cluster_id"), row.get("cluster") or "Sin clúster")
        bucket = heat.setdefault(
            key,
            {
                "cluster": key[1],
                "cluster_id": key[0],
                "year1": 0.0,
                "year2": 0.0,
                "year3": 0.0,
                "n": 0,
                "n_with_budget": 0,
            },
        )
        bucket["year1"] += float(row.get("year1") or 0)
        bucket["year2"] += float(row.get("year2") or 0)
        bucket["year3"] += float(row.get("year3") or 0)
        bucket["n"] += 1
        if row.get("year1") or row.get("year2") or row.get("year3"):
            bucket["n_with_budget"] += 1
    heatmap = sorted(
        heat.values(),
        key=lambda b: (-(b["year1"] + b["year2"] + b["year3"]), b["cluster"]),
    )
    items = [
        {
            "cluster": row.get("cluster"),
            "year1": float(row.get("year1") or 0),
            "year2": float(row.get("year2") or 0),
            "year3": float(row.get("year3") or 0),
            "points": int(row.get("points") or 0),
            "technology_id": row.get("technology_id"),
            "name": row.get("name"),
        }
        for row in rows
    ]
    comparators = [
        {"technology_id": row["technology_id"], "name": row.get("name"), "comparators": row["comparators"]}
        for row in rows
        if row.get("comparators")
    ]
    return {
        "budget_heatmap": heatmap,
        "budget_items": items,
        "budget_unparsed": sum(1 for row in rows if row.get("budget_unparsed")),
        "comparators": comparators,
    }


def build_payload(
    db: Session,
    cycle: Cycle,
    *,
    cluster_id: int | None = None,
    tech_type_id: int | None = None,
    band: str = "",
    status: str = "",
    phase: str = "",
    date_from: date | None = None,
    date_to: date | None = None,
    priority_min: int | None = None,
    stage: str = "",
    rows: list[dict] | None = None,
) -> dict:
    """Payload completo (incluye capa restringida). No se expone tal cual por API."""
    all_rows = rows if rows is not None else collect_rows(db, cycle)
    filters = normalize_filters(
        cluster_id=cluster_id,
        tech_type_id=tech_type_id,
        band=band,
        status=status,
        phase=phase,
        date_from=date_from,
        date_to=date_to,
        priority_min=priority_min,
        stage=stage,
    )
    selected = filter_rows(all_rows, filters)
    agg = aggregate(selected)
    budget = _budget(selected)
    return {
        "schema": DATAMART_SCHEMA,
        "cycle_id": cycle.id,
        "cycle_code": cycle.code,
        **agg,
        "kpis": {**agg["funnel"], **agg["conversion"]},
        "ttm_scatter": [_public_row(r) for r in selected],
        **budget,
        "total_in_cycle": len(all_rows),
    }


def refresh_datamart(db: Session, cycle: Cycle) -> CycleDatamart:
    rows = collect_rows(db, cycle)
    payload = build_payload(db, cycle, rows=rows)
    payload["rows"] = rows
    payload["ttm_thresholds"] = ttm.load_params(db)
    row = db.query(CycleDatamart).filter(CycleDatamart.cycle_id == cycle.id).first()
    if row is None:
        row = CycleDatamart(cycle_id=cycle.id, payload=payload)
        db.add(row)
    else:
        row.payload = payload
    row.refreshed_at = _utcnow()
    return row


def _cached_rows(cached: CycleDatamart | None) -> list[dict] | None:
    payload = (cached.payload or {}) if cached else {}
    if int(payload.get("schema") or 0) >= DATAMART_SCHEMA and isinstance(payload.get("rows"), list):
        return payload["rows"]
    return None


def dashboard_for(
    db: Session,
    cycle: Cycle,
    *,
    include_restricted: bool,
    cluster_id: int | None = None,
    tech_type_id: int | None = None,
    band: str = "",
    status: str = "",
    phase: str = "",
    date_from: date | None = None,
    date_to: date | None = None,
    priority_min: int | None = None,
    stage: str = "",
) -> dict:
    """Tablero del ciclo con la capa de acceso aplicada.

    Ciclo cerrado con datamart vigente: todo (tambien lo filtrado) se resuelve
    sobre las filas del datamart. Ciclo abierto: se calcula en vivo.
    """
    filters = normalize_filters(
        cluster_id=cluster_id,
        tech_type_id=tech_type_id,
        band=band,
        status=status,
        phase=phase,
        date_from=date_from,
        date_to=date_to,
        priority_min=priority_min,
        stage=stage,
    )
    cached = db.query(CycleDatamart).filter(CycleDatamart.cycle_id == cycle.id).first()
    rows = _cached_rows(cached) if cycle.status == "cerrado_consolidado" else None
    use_cache = rows is not None
    if rows is None:
        rows = collect_rows(db, cycle)
    thresholds = (
        (cached.payload or {}).get("ttm_thresholds") if use_cache and cached else None
    ) or ttm.load_params(db)
    payload = build_payload(db, cycle, rows=rows, **filters)

    active = _active(filters)
    public = {
        "cycle_id": cycle.id,
        "cycle_code": cycle.code,
        "funnel": payload["funnel"],
        "kpis": payload["kpis"],
        "conversion": payload["conversion"],
        "by_cluster": payload["by_cluster"],
        "by_type": payload["by_type"],
        "by_band": payload["by_band"],
        "by_phase": payload["by_phase"],
        "ttm_scatter": payload["ttm_scatter"],
        "ttm_summary": payload["ttm_summary"],
        "ttm_thresholds": thresholds,
        "total_in_cycle": payload["total_in_cycle"],
        "filters_applied": {k: (v.isoformat() if isinstance(v, date) else v) for k, v in active.items()},
        "restricted": include_restricted,
        "from_cache": use_cache,
        "refreshed_at": cached.refreshed_at if (cached and use_cache) else None,
        "cycle_status": cycle.status,
    }
    if include_restricted:
        public["budget_heatmap"] = payload["budget_heatmap"]
        public["budget_items"] = payload["budget_items"]
        public["budget_unparsed"] = payload["budget_unparsed"]
        public["comparators"] = payload["comparators"]
    return public


def _comparators(db: Session, cycle: Cycle) -> list[dict]:
    return _budget(collect_rows(db, cycle))["comparators"]


# --------------------------------------------------------------------------- #
#  Exportacion del recorte
# --------------------------------------------------------------------------- #
EXPORT_COLUMNS = (
    ("technology_id", "ID"),
    ("name", "Tecnologia"),
    ("cluster", "Cluster"),
    ("tech_type", "Tipologia"),
    ("status_label", "Estado en el ciclo"),
    ("stage_label", "Etapa alcanzada del embudo"),
    ("points", "Puntos P1-P6"),
    ("priority_pct", "%P"),
    ("phase", "Fase declarada"),
    ("phase_label", "Fase clinica (grupo)"),
    ("band_label", "Franja time-to-market"),
    ("months", "Meses al mercado"),
    ("basis_label", "Base del calculo TTM"),
    ("reference_date", "Fecha de referencia TTM"),
    ("ttm_note", "Nota TTM"),
    ("captured_on", "Fecha de captura"),
)
EXPORT_RESTRICTED_COLUMNS = (
    ("year1", "Impacto presupuestal anio 1"),
    ("year2", "Impacto presupuestal anio 2"),
    ("year3", "Impacto presupuestal anio 3"),
)


def _ttm_note(row: dict) -> str:
    if row.get("months") is None:
        return "Sin dato: falta fecha de fin de fase III o de aprobación"
    if row.get("basis") == "registro_invima":
        return "Ya tiene registro INVIMA"
    if row.get("approved"):
        return "Ya aprobada por FDA/EMA"
    if row.get("overdue"):
        return "La fecha esperada ya pasó sin aprobación registrada"
    return ""


def export_table(data: dict, *, include_restricted: bool) -> tuple[list[str], list[list]]:
    """Encabezados y filas del recorte, con los tipos nativos (numeros como numeros)."""
    columns = list(EXPORT_COLUMNS)
    if include_restricted:
        columns += list(EXPORT_RESTRICTED_COLUMNS)
    budget = {b["technology_id"]: b for b in (data.get("budget_items") or [])} if include_restricted else {}
    table: list[list] = []
    for row in data.get("ttm_scatter") or []:
        full = dict(row)
        full["stage_label"] = FUNNEL_LABELS.get(row.get("stage"), row.get("stage") or "")
        full["points"] = int(row["points"]) if row.get("points_rated") else None
        full["ttm_note"] = _ttm_note(row)
        if include_restricted:
            item = budget.get(row["technology_id"]) or {}
            for key in ("year1", "year2", "year3"):
                full[key] = float(item.get(key) or 0) or None
        table.append([full.get(key) for key, _ in columns])
    return [label for _, label in columns], table


def _csv_cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "Si" if value else "No"
    if isinstance(value, float):
        text = f"{value:.2f}".rstrip("0").rstrip(".") if value % 1 else f"{value:.0f}"
        return text.replace(".", ",")
    text = str(value)
    # Neutraliza formulas al abrir el CSV en una hoja de calculo.
    if text[:1] in ("=", "+", "-", "@", "\t", "\r"):
        text = "'" + text
    return text


def export_dashboard(cycle: Cycle, data: dict, *, fmt: str, include_restricted: bool) -> tuple[bytes, str, str]:
    headers, table = export_table(data, include_restricted=include_restricted)
    stamp = ttm.today_co().isoformat()
    slug = re.sub(r"[^A-Za-z0-9]+", "-", cycle.code or f"ciclo-{cycle.id}").strip("-").lower()
    base = f"tablero-{slug}-{stamp}"
    filters = data.get("filters_applied") or {}
    filters_text = ", ".join(f"{k}={v}" for k, v in sorted(filters.items())) or "sin filtros"
    if fmt == "xlsx":
        from io import BytesIO

        from openpyxl import Workbook
        from openpyxl.styles import Font
        from openpyxl.utils import get_column_letter

        wb = Workbook()
        ws = wb.active
        ws.title = "Recorte"
        ws.append(headers)
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for values in table:
            ws.append(
                [
                    ("'" + v) if isinstance(v, str) and v[:1] in ("=", "+", "-", "@") else v
                    for v in values
                ]
            )
        for idx, header in enumerate(headers, start=1):
            width = max([len(str(header))] + [len(str(r[idx - 1] or "")) for r in table[:200]])
            ws.column_dimensions[get_column_letter(idx)].width = min(60, max(10, width + 2))
        ws.freeze_panes = "A2"
        meta = wb.create_sheet("Contexto")
        meta.append(["Ciclo", cycle.code])
        meta.append(["Estado del ciclo", data.get("cycle_status") or cycle.status])
        meta.append(["Filtros", filters_text])
        meta.append(["Tecnologias en el recorte", len(table)])
        meta.append(["Fuente", "Datamart al cierre" if data.get("from_cache") else "Calculo en vivo"])
        meta.append(["Capa", "Restringida (con montos)" if include_restricted else "Agregada (sin montos)"])
        meta.append(["Generado", datetime.now(COLOMBIA_TZ).strftime("%Y-%m-%d %H:%M")])
        buffer = BytesIO()
        wb.save(buffer)
        return (
            buffer.getvalue(),
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            f"{base}.xlsx",
        )

    import csv
    import io

    out = io.StringIO()
    writer = csv.writer(out, delimiter=";", lineterminator="\r\n")
    writer.writerow(headers)
    for values in table:
        writer.writerow([_csv_cell(v) for v in values])
    return ("﻿" + out.getvalue()).encode("utf-8"), "text/csv; charset=utf-8", f"{base}.csv"


def _public_doc_filter():
    """Criterio unico de lo publico: expediente publicado y no confidencial.

    Lo usan la lista publica (`list_public_fiches`) y las estadisticas de
    transparencia (`public_stats`), de modo que el conteo nunca supera a las
    fichas visibles ni revela que existen publicaciones confidenciales.
    """
    return (EvaluationDoc.status == "publicado", EvaluationDoc.confidential == False)  # noqa: E712


def public_stats(db: Session) -> dict:
    public_ids = (
        db.query(EvaluationDoc.technology_id)
        .filter(*_public_doc_filter())
        .distinct()
        .subquery()
    )
    published = db.query(func.count()).select_from(public_ids).scalar() or 0
    cycles = db.query(func.count(Cycle.id)).filter(Cycle.is_historic == False).scalar() or 0  # noqa: E712
    by_cluster = (
        db.query(Cluster.name, func.count(Technology.id))
        .select_from(Technology)
        .join(public_ids, public_ids.c.technology_id == Technology.id)
        .outerjoin(Cluster, Cluster.id == Technology.cluster_id)
        .group_by(Cluster.name)
        .all()
    )
    items = [{"label": name or "Sin clúster", "value": int(n)} for name, n in by_cluster]
    return {
        "published": int(published),
        "cycles": cycles,
        "by_cluster": sorted(items, key=lambda i: (-i["value"], i["label"])),
    }


def sanitize_public_body(doc: EvaluationDoc) -> dict:
    hidden = set(PUBLIC_HIDDEN_FIELDS) | set(doc.confidential_fields or [])
    return {k: v for k, v in (doc.body or {}).items() if k not in hidden and str(v or "").strip()}


def list_public_fiches(
    db: Session, *, q: str = "", cluster_id: int | None = None, limit: int = 50
) -> list[dict]:
    query = (
        db.query(EvaluationDoc, Technology)
        .join(Technology, Technology.id == EvaluationDoc.technology_id)
        .filter(*_public_doc_filter())
    )
    if cluster_id:
        query = query.filter(Technology.cluster_id == cluster_id)
    q = (q or "").strip()[:80]
    if q:
        like = f"%{q}%"
        query = query.filter(
            or_(
                Technology.commercial_name.ilike(like),
                Technology.inn_name.ilike(like),
                EvaluationDoc.title.ilike(like),
            )
        )
    items = []
    cap = max(1, min(int(limit or 50), 50))
    seen: set[int] = set()
    for doc, tech in query.order_by(EvaluationDoc.published_at.desc()).all():
        if tech.id in seen:
            continue
        seen.add(tech.id)
        items.append(
            {
                "id": tech.id,
                "doc_id": doc.id,
                "title": doc.title or _tech_name(tech),
                "commercial_name": tech.commercial_name,
                "inn_name": tech.inn_name,
                "nct_ids": list(tech.nct_ids or []),
                "product_level": doc.product_level,
                "published_at": doc.published_at,
            }
        )
        if len(items) >= cap:
            break
    return items


def render_public_fiche_html(data: dict) -> str:
    from .report_html import render_public_fiche_html as _render

    return _render(data)


def public_fiche(db: Session, technology_id: int) -> dict:
    doc = (
        db.query(EvaluationDoc)
        .filter(
            EvaluationDoc.technology_id == technology_id,
            EvaluationDoc.status == "publicado",
            EvaluationDoc.confidential == False,  # noqa: E712
        )
        .order_by(EvaluationDoc.published_at.desc())
        .first()
    )
    tech = db.get(Technology, technology_id)
    if doc is None or tech is None:
        raise StrategyRuleError("Ficha pública no disponible.")
    calc = ttm.compute_for(db, tech)
    cluster = db.get(Cluster, tech.cluster_id) if tech.cluster_id else None
    from .evaluation import FIELD_LABELS

    body = sanitize_public_body(doc)
    return {
        "id": tech.id,
        "doc_id": doc.id,
        "title": doc.title or _tech_name(tech),
        "commercial_name": tech.commercial_name,
        "inn_name": tech.inn_name,
        "manufacturer": tech.manufacturer,
        "nct_ids": list(tech.nct_ids or []),
        "cluster": cluster.name if cluster else "",
        "product_level": doc.product_level,
        "published_at": doc.published_at,
        "ttm_band": calc["band"],
        "ttm_band_label": calc["band_label"],
        "body": body,
        "field_labels": {k: FIELD_LABELS.get(k, k) for k in body},
    }


def compile_bulletin(db: Session, cycle: Cycle, *, actor: str = "") -> Bulletin:
    payload = build_payload(db, cycle)
    row = (
        db.query(Bulletin)
        .filter(Bulletin.cycle_id == cycle.id, Bulletin.status != "publicado")
        .order_by(Bulletin.id.desc())
        .first()
    )
    if row is None:
        row = Bulletin(cycle_id=cycle.id)
        db.add(row)
    row.title = f"Boletín epidemiológico y financiero — {cycle.code}"
    row.status = "pendiente_aprobacion"
    # Recompilar cambia el contenido: la aprobacion previa correspondia a otras
    # cifras y no puede habilitar la publicacion de las nuevas (RF19).
    row.approved_by = ""
    row.body = {
        "cycle_code": cycle.code,
        "funnel": payload["funnel"],
        "by_cluster": payload["by_cluster"],
        "by_band": [b for b in payload["by_band"] if b.get("value")],
        "prioritized": payload["funnel"].get("prioritized", 0),
        "published": payload["funnel"].get("published", 0),
        "compiled_on": date.today().isoformat(),
    }
    row.compiled_by = actor
    db.flush()
    audit.record_action(
        db,
        entity_type="bulletins",
        entity_id=row.id,
        action="bulletin:compile",
        new_value={"cycle_id": cycle.id, "status": row.status},
    )
    return row


def approve_bulletin(db: Session, bulletin: Bulletin, *, actor: str, publish: bool = False) -> Bulletin:
    if bulletin.status == "publicado":
        raise StrategyRuleError("El boletín ya está publicado.")
    bulletin.approved_by = actor
    if publish:
        bulletin.status = "publicado"
        bulletin.published_at = _utcnow()
    else:
        bulletin.status = "pendiente_aprobacion"
    audit.record_action(
        db,
        entity_type="bulletins",
        entity_id=bulletin.id,
        action="bulletin:approve" if not publish else "bulletin:publish",
        new_value={"status": bulletin.status, "approved_by": actor},
    )
    if publish and not bulletin.approved_by:
        raise StrategyRuleError("El boletín requiere aprobación del líder antes de publicarse.")
    return bulletin


def publish_bulletin(db: Session, bulletin: Bulletin, *, actor: str) -> Bulletin:
    if not (bulletin.approved_by or "").strip():
        raise StrategyRuleError("El boletín requiere aprobación del líder antes de publicarse.")
    bulletin.status = "publicado"
    bulletin.published_at = _utcnow()
    audit.record_action(
        db,
        entity_type="bulletins",
        entity_id=bulletin.id,
        action="bulletin:publish",
        new_value={"approved_by": bulletin.approved_by},
    )
    return bulletin


def render_bulletin_html(bulletin: Bulletin) -> str:
    from .report_html import render_bulletin_html as _render

    return _render(bulletin)


def on_cycle_closed(db: Session, cycle: Cycle, *, actor: str = "") -> None:
    snapshot_ttm(db, cycle)
    refresh_datamart(db, cycle)
    compile_bulletin(db, cycle, actor=actor)


def _subscribers(db: Session, cluster_id: int | None) -> list[AlertSubscription]:
    rows = db.query(AlertSubscription).filter(AlertSubscription.enabled == True).all()  # noqa: E712
    return [s for s in rows if s.cluster_id is None or s.cluster_id == cluster_id]


def _notify(db: Session, *, kind: str, title: str, body: str, tech: Technology) -> int:
    sent = 0
    for sub in _subscribers(db, tech.cluster_id):
        exists = (
            db.query(AlertEvent)
            .filter(
                AlertEvent.user_id == sub.user_id,
                AlertEvent.kind == kind,
                AlertEvent.technology_id == tech.id,
                AlertEvent.read_at.is_(None),
            )
            .first()
        )
        if exists:
            continue
        db.add(
            AlertEvent(
                user_id=sub.user_id,
                kind=kind,
                title=title,
                body=body,
                technology_id=tech.id,
                cluster_id=tech.cluster_id,
            )
        )
        sent += 1
        # P4-2: con SMTP configurado la alerta tambien sale por correo (en segundo
        # plano, sin demorar ni romper la operacion que la disparo).
        user = db.get(User, sub.user_id)
        if user is not None and user.is_active and user.email:
            from . import mailer

            base = (get_settings_public_base_url() or "").rstrip("/")
            mailer.alert_notification(
                to=user.email, title=title, body=body, link=f"{base}/alertas" if base else ""
            )
    return sent


def get_settings_public_base_url() -> str:
    from .config import settings

    return settings.public_base_url or ""


def watch_technology(db: Session, tech: Technology, *, previous_phase: str = "", previous_status: str = "") -> None:
    """Dispara alertas de fase III en pais y de cambio en alto riesgo presupuestal."""
    threshold = int(get_param(db, "alerts.high_budget_points", 5) or 5)
    points = 0
    latest = (
        db.query(CycleTechnology)
        .filter(CycleTechnology.technology_id == tech.id)
        .order_by(CycleTechnology.id.desc())
        .first()
    )
    if latest and latest.priority_points is not None:
        points = int(latest.priority_points)

    now_phase = tech.development_phase or ""
    if _looks_phase3(now_phase) and not _looks_phase3(previous_phase) and _mentions_colombia(db, tech):
        _notify(
            db,
            kind="new_phase3_colombia",
            title=f"Nuevo ensayo fase III en el país: {_tech_name(tech)}",
            body="Se detectó una señal de ensayo fase III asociada a Colombia.",
            tech=tech,
        )

    if points >= threshold and previous_status and previous_status != tech.status:
        _notify(
            db,
            kind="phase_change_high_budget",
            title=f"Cambio de fase en tecnología de alto riesgo: {_tech_name(tech)}",
            body=f"Paso de {previous_status} a {tech.status} con {points} puntos de priorización.",
            tech=tech,
        )


def subscribe(db: Session, user: User, cluster_id: int | None) -> AlertSubscription:
    row = (
        db.query(AlertSubscription)
        .filter(
            AlertSubscription.user_id == user.id,
            AlertSubscription.cluster_id == cluster_id
            if cluster_id is not None
            else AlertSubscription.cluster_id.is_(None),
        )
        .first()
    )
    if row:
        row.enabled = True
        return row
    row = AlertSubscription(user_id=user.id, cluster_id=cluster_id, enabled=True)
    db.add(row)
    db.flush()
    return row
