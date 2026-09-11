"""Frente C: precision numerica del tablero estrategico (RF17) y de la bandeja.

Cada agregado se compara contra conteos calculados aqui de forma independiente
(a partir de la definicion de la regla, no llamando a las funciones del
servicio) sobre datos sembrados con semilla fija. Se recorren cientos de
combinaciones de filtros y se exige coincidencia exacta.
"""
from __future__ import annotations

import csv
import datetime as dt
import io
import itertools
import random
import statistics
import sys
import time
from collections import Counter
from pathlib import Path

import pytest
from fastapi import Depends, FastAPI
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import Session, sessionmaker
from sqlalchemy.pool import StaticPool

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from app import cycle_service, graph_service, methodology, rbac, strategy_service, ttm  # noqa: E402
from app.database import Base, get_db  # noqa: E402
from app.deps import get_current_user  # noqa: E402
from app.models import (  # noqa: E402
    AlertEvent,
    Bulletin,
    Cluster,
    Cycle,
    CycleDatamart,
    CycleTechnology,
    EvaluationDoc,
    Finding,
    MergeProposal,
    MethodologyParam,
    NoveltyAssessment,
    PriorityScore,
    ReviewAssignment,
    Source,
    Submission,
    TechType,
    Technology,
    TimeToMarketSnapshot,
    User,
)
from app.routers import dashboard as dashboard_router  # noqa: E402
from app.routers import strategy as strategy_router  # noqa: E402

# --------------------------------------------------------------------------- #
#  Reglas de referencia (escritas desde la especificacion, no desde el codigo)
# --------------------------------------------------------------------------- #
FILTERED = {"filtrada_apta_priorizacion", "priorizada", "bajo_vigilancia", "no_priorizada", "en_evaluacion", "publicada"}
PRIORITIZED = {"priorizada", "en_evaluacion", "publicada"}
STAGES = ["captured", "filtered", "prioritized", "evaluated", "published"]
ALL_STATUSES = [
    "asignada_a_ciclo",
    "filtrada_apta_priorizacion",
    "excluida",
    "priorizada",
    "bajo_vigilancia",
    "no_priorizada",
    "en_evaluacion",
    "publicada",
]
# Texto de fase declarado -> grupo esperado.
PHASE_CASES = [
    ("Fase III", "fase_iii"),
    ("Phase 3", "fase_iii"),
    ("PHASE2", "fase_ii"),
    ("PHASE1|PHASE2", "fase_ii"),
    ("Fase II/III", "fase_iii"),
    ("Fase I, Fase II", "fase_ii"),
    ("Early Phase 1", "fase_i"),
    ("EARLY_PHASE1", "fase_i"),
    ("Fase IV", "fase_iv"),
    ("Autorizado / extensiones fase III", "autorizada"),
    ("Aprobado", "autorizada"),
    ("Approved", "autorizada"),
    ("No aprobado", "sin_fase"),
    ("", "sin_fase"),
    ("Evidencia publicada", "sin_fase"),
    ("Fase 1", "fase_i"),
    ("fase iii", "fase_iii"),
    ("NA", "sin_fase"),
]
BUDGET_TEXTS = [
    ("", 0.0),
    ("3200000000. Anio 1: 40-80 pacientes", 3.2e9),
    ("COP 3.200.000.000", 3.2e9),
    ("1.200 millones", 1.2e9),
    ("3,5 mil millones", 3.5e9),
    ("por estimar con precio de referencia", 0.0),
]


def ref_stage(status: str, doc_status: str | None) -> str:
    if status not in FILTERED:
        return "captured"
    if status not in PRIORITIZED:
        return "filtered"
    evaluated = status in {"en_evaluacion", "publicada"} or doc_status is not None
    if not evaluated:
        return "prioritized"
    if status == "publicada" or doc_status == "publicado":
        return "published"
    return "evaluated"


def ref_ttm(tech_data: dict, as_of: dt.date, *, review_days=180, inm=12, tra=24, eme=36):
    approvals = [d for d in (tech_data["fda"], tech_data["ema"]) if d]
    if approvals:
        ref = min(approvals)
    elif tech_data["invima"]:
        return 0.0, "inminente"
    elif tech_data["phase3"]:
        ref = tech_data["phase3"] + dt.timedelta(days=review_days)
    else:
        return None, "desconocido"
    months = max(round((ref - as_of).days / 30.44, 2), 0.0)
    if months <= inm:
        band = "inminente"
    elif months <= tra:
        band = "transicion"
    elif months <= eme:
        band = "emergente"
    else:
        band = "lejano"
    return months, band


# --------------------------------------------------------------------------- #
#  Arnes: base en memoria compartida + app minima con los routers del frente C
# --------------------------------------------------------------------------- #
class Harness:
    def __init__(self):
        self.engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
        Base.metadata.create_all(bind=self.engine)
        self.Session = sessionmaker(bind=self.engine, autoflush=False)
        self.db: Session = self.Session()
        methodology.seed_catalogs(self.db)
        self.users: dict[str, User] = {}
        for role in rbac.ROLES:
            user = User(email=f"c.{role}@iets.org.co", name=f"Usuario {role}", role=role, is_active=True)
            self.db.add(user)
            self.users[role] = user
        self.db.commit()
        self.current = "superadmin"
        app = FastAPI()
        app.include_router(strategy_router.router)
        app.include_router(dashboard_router.router)

        def _get_db():
            session = self.Session()
            try:
                yield session
            finally:
                session.close()

        def _user(db: Session = Depends(get_db)) -> User:
            return db.query(User).filter(User.email == f"c.{self.current}@iets.org.co").first()

        app.dependency_overrides[get_db] = _get_db
        app.dependency_overrides[get_current_user] = _user
        self.client = TestClient(app)

    def as_role(self, role: str) -> TestClient:
        self.current = role
        return self.client


@pytest.fixture()
def h():
    harness = Harness()
    yield harness
    harness.db.close()


def _cycle(db, code, status="en_priorizacion", **kw) -> Cycle:
    today = ttm.today_co()
    cycle = Cycle(
        code=code,
        year=today.year,
        opened_on=kw.pop("opened_on", today - dt.timedelta(weeks=4)),
        data_cutoff_on=kw.pop("data_cutoff_on", today + dt.timedelta(weeks=6)),
        bulletin_due_on=kw.pop("bulletin_due_on", today + dt.timedelta(weeks=8)),
        status=status,
        **kw,
    )
    db.add(cycle)
    db.commit()
    return cycle


def seed_cycle(db: Session, cycle: Cycle, n: int, seed: int, *, as_of: dt.date) -> list[dict]:
    """Siembra n tecnologias con atributos aleatorios y devuelve su verdad de referencia."""
    rng = random.Random(seed)
    clusters = db.query(Cluster).order_by(Cluster.id).all()
    types = db.query(TechType).order_by(TechType.id).all()
    truth = []
    for i in range(n):
        cluster = rng.choice(clusters + [None])
        ttype = rng.choice(types + [None])
        phase_text, phase_bucket = rng.choice(PHASE_CASES)
        kind = rng.random()
        fda = ema = phase3 = None
        invima = ""
        if kind < 0.2:
            fda = as_of + dt.timedelta(days=rng.randint(-2000, 1500))
        elif kind < 0.3:
            ema = as_of + dt.timedelta(days=rng.randint(-900, 900))
            fda = rng.choice([None, as_of + dt.timedelta(days=rng.randint(-900, 900))])
        elif kind < 0.35:
            invima = f"INVIMA 2024M-{i:05d}"
        elif kind < 0.8:
            phase3 = as_of + dt.timedelta(days=rng.randint(-1500, 1800))
        # Hora UTC variada: cerca de medianoche UTC la fecha local de Colombia es la del dia anterior.
        captured_at = dt.datetime(2026, 1, 1, tzinfo=dt.timezone.utc) + dt.timedelta(
            days=rng.randint(0, 240), hours=rng.choice([0, 2, 4, 5, 6, 12, 23])
        )
        tech = Technology(
            commercial_name=f"{cycle.code} T{i:04d} " + "nombre largo " * rng.randint(0, 12),
            inn_name=f"inn{i}",
            cluster_id=cluster.id if cluster else None,
            tech_type_id=ttype.id if ttype else None,
            development_phase=phase_text,
            fda_approval_date=fda,
            ema_approval_date=ema,
            phase3_completion_date=phase3,
            invima_registry=invima,
            captured_at=captured_at,
            status="asignada_a_ciclo",
        )
        db.add(tech)
        db.flush()
        status = rng.choice(ALL_STATUSES)
        points = rng.choice([None, 0, 1, 2, 3, 4, 5, 6]) if status != "asignada_a_ciclo" else None
        db.add(
            CycleTechnology(
                cycle_id=cycle.id,
                technology_id=tech.id,
                status=status,
                priority_points=points,
                priority_pct=round(points / 6 * 100, 2) if points is not None else None,
            )
        )
        doc_status = None
        budget = [0.0, 0.0, 0.0]
        if status in PRIORITIZED and rng.random() < 0.75:
            doc_status = rng.choice(["borrador", "revision_interna", "con_observaciones", "aprobado_comite", "publicado"])
            texts = [rng.choice(BUDGET_TEXTS) for _ in range(3)]
            budget = [value for _, value in texts]
            db.add(
                EvaluationDoc(
                    cycle_id=cycle.id,
                    technology_id=tech.id,
                    status=doc_status,
                    product_level="mini_hta",
                    title=f"Doc {i}",
                    body={
                        "budget_year_1": texts[0][0],
                        "budget_year_2": texts[1][0],
                        "budget_year_3": texts[2][0],
                        "comparators_sgsss": "Comparador" if rng.random() < 0.5 else "",
                    },
                )
            )
        months, band = ref_ttm({"fda": fda, "ema": ema, "invima": invima, "phase3": phase3}, as_of)
        truth.append(
            {
                "technology_id": tech.id,
                "cluster_id": cluster.id if cluster else None,
                "cluster": cluster.name if cluster else "Sin clúster",
                "tech_type_id": ttype.id if ttype else None,
                "tech_type": ttype.name if ttype else "Sin tipología",
                "status": status,
                "points": points,
                "phase": phase_bucket,
                "months": months,
                "band": band,
                "captured_on": (captured_at - dt.timedelta(hours=5)).date(),
                "stage": ref_stage(status, doc_status),
                "budget": budget,
            }
        )
    db.commit()
    return truth


def ref_filter(truth, f):
    out = []
    for t in truth:
        if f.get("cluster_id") is not None:
            want = None if f["cluster_id"] == 0 else f["cluster_id"]
            if t["cluster_id"] != want:
                continue
        if f.get("tech_type_id") is not None:
            want = None if f["tech_type_id"] == 0 else f["tech_type_id"]
            if t["tech_type_id"] != want:
                continue
        if f.get("band") and t["band"] != f["band"]:
            continue
        if f.get("phase") and t["phase"] != f["phase"]:
            continue
        if f.get("status") and t["status"] != f["status"]:
            continue
        if f.get("stage") and STAGES.index(t["stage"]) < STAGES.index(f["stage"]):
            continue
        if f.get("priority_min") is not None and (t["points"] is None or t["points"] < f["priority_min"]):
            continue
        if f.get("date_from") and t["captured_on"] < f["date_from"]:
            continue
        if f.get("date_to") and t["captured_on"] > f["date_to"]:
            continue
        out.append(t)
    return out


def ref_aggregates(rows):
    funnel = {stage: sum(1 for r in rows if STAGES.index(r["stage"]) >= i) for i, stage in enumerate(STAGES)}
    return {
        "funnel": funnel,
        "by_cluster": Counter(r["cluster"] for r in rows),
        "by_type": Counter(r["tech_type"] for r in rows),
        "by_band": Counter(r["band"] for r in rows),
        "by_phase": Counter(r["phase"] for r in rows),
        "ids": sorted(r["technology_id"] for r in rows),
        "budget": [round(sum(r["budget"][k] for r in rows), 2) for k in range(3)],
        "without": sum(1 for r in rows if r["months"] is None),
    }


def assert_matches(data: dict, expected: dict, context: str):
    assert data["funnel"] == expected["funnel"], context
    assert {i["label"]: i["value"] for i in data["by_cluster"]} == dict(expected["by_cluster"]), context
    assert {i["label"]: i["value"] for i in data["by_type"]} == dict(expected["by_type"]), context
    bands = {i["code"]: i["value"] for i in data["by_band"]}
    assert set(bands) == set(ttm.TTM_BANDS), context
    assert {k: v for k, v in bands.items() if v} == dict(expected["by_band"]), context
    phases = {i["code"]: i["value"] for i in data["by_phase"]}
    assert {k: v for k, v in phases.items() if v} == dict(expected["by_phase"]), context
    assert sorted(r["technology_id"] for r in data["ttm_scatter"]) == expected["ids"], context
    assert data["ttm_summary"]["without_estimate"] == expected["without"], context
    assert sum(i["value"] for i in data["by_cluster"]) == data["funnel"]["captured"], context
    if data.get("budget_heatmap") is not None:
        got = [round(sum(b[k] for b in data["budget_heatmap"]), 2) for k in ("year1", "year2", "year3")]
        assert got == expected["budget"], context


# --------------------------------------------------------------------------- #
#  1. Motor de time-to-market: regla ttm.* y bordes de los umbrales
# --------------------------------------------------------------------------- #
@pytest.fixture()
def db():
    engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    Base.metadata.create_all(bind=engine)
    session = sessionmaker(bind=engine, autoflush=False)()
    methodology.seed_catalogs(session)
    yield session
    session.close()


AS_OF = dt.date(2026, 9, 10)


@pytest.mark.parametrize(
    "days,band",
    [
        (0, "inminente"),
        (365, "inminente"),  # 11,99 meses
        (366, "transicion"),  # 12,02
        (730, "transicion"),  # 23,98
        (731, "emergente"),  # 24,01
        (1095, "emergente"),  # 35,97
        (1097, "lejano"),  # 36,04
    ],
)
def test_ttm_band_borders_follow_parameters(db, days, band):
    tech = Technology(fda_approval_date=AS_OF + dt.timedelta(days=days))
    calc = ttm.compute_for(db, tech, as_of=AS_OF)
    months = round(days / 30.44, 2)
    assert calc["months"] == months
    expected = "inminente" if months <= 12 else "transicion" if months <= 24 else "emergente" if months <= 36 else "lejano"
    assert calc["band"] == expected == band


def test_ttm_phase3_plus_review_days_and_zero_is_respected(db):
    tech = Technology(phase3_completion_date=AS_OF)
    assert ttm.compute_for(db, tech, as_of=AS_OF)["months"] == round(180 / 30.44, 2)
    # Bug corregido: un plazo de 0 dias se confundia con "sin configurar" y volvia a 180.
    db.get(MethodologyParam, "ttm.regulatory_review_days").value = "0"
    db.commit()
    calc = ttm.compute_for(db, tech, as_of=AS_OF)
    assert calc["months"] == 0.0
    assert calc["overdue"] is False
    # Umbral inminente en 0: solo lo que ya llego es inminente.
    db.get(MethodologyParam, "ttm.inminente_months_max").value = "0"
    db.commit()
    assert ttm.classify_months(db, 0.0) == "inminente"
    assert ttm.classify_months(db, 0.01) == "transicion"


def test_ttm_non_monotonic_thresholds_are_made_contiguous(db):
    db.get(MethodologyParam, "ttm.inminente_months_max").value = "30"
    db.get(MethodologyParam, "ttm.transicion_months_max").value = "10"
    db.commit()
    params = ttm.load_params(db)
    assert params["transicion"] == 30 and params["emergente"] >= 30
    assert ttm.classify_months(db, 25) == "inminente"
    assert ttm.classify_months(db, 31) in ("emergente", "lejano")


def test_ttm_distinguishes_no_data_from_zero_months(db):
    none = ttm.compute_for(db, Technology(), as_of=AS_OF)
    assert none["months"] is None and none["months_raw"] is None
    assert none["band"] == "desconocido" and none["basis"] == "sin_fase_iii"

    approved = ttm.compute_for(db, Technology(fda_approval_date=dt.date(2018, 1, 26), ema_approval_date=dt.date(2017, 9, 26)), as_of=AS_OF)
    assert approved["months"] == 0.0 and approved["approved"] is True and approved["overdue"] is False
    assert approved["reference_date"] == "2017-09-26"  # la primera aprobacion
    assert approved["months_raw"] == round((dt.date(2017, 9, 26) - AS_OF).days / 30.44, 2) < 0

    overdue = ttm.compute_for(db, Technology(phase3_completion_date=dt.date(2022, 6, 30)), as_of=AS_OF)
    assert overdue["months"] == 0.0 and overdue["overdue"] is True and overdue["approved"] is False
    assert overdue["reference_date"] == "2022-12-27"

    registry = ttm.compute_for(db, Technology(invima_registry="INVIMA 2020M-1"), as_of=AS_OF)
    assert registry["months"] == 0.0 and registry["basis"] == "registro_invima" and registry["approved"] is True

    # Un datetime en un campo de fecha ya no rompe el calculo (datetime - date).
    stamped = ttm.compute_for(db, Technology(fda_approval_date=dt.datetime(2027, 3, 10, 15, 0)), as_of=AS_OF)
    assert stamped["months"] == round((dt.date(2027, 3, 10) - AS_OF).days / 30.44, 2)


@pytest.mark.parametrize("text,bucket", PHASE_CASES)
def test_phase_bucket(text, bucket):
    assert strategy_service._phase_bucket(text) == bucket


@pytest.mark.parametrize(
    "text,amount",
    [
        ("3200000000. Anio 1: 40-80 pacientes", 3.2e9),
        ("COP 3.200.000.000", 3.2e9),
        ("$3,200,000,000", 3.2e9),
        ("3.200 millones", 3.2e9),
        ("3,5 mil millones", 3.5e9),
        ("Anio 1: 40-80 pacientes; 1.200 millones de pesos", 1.2e9),
        ("Anio 2026: 750 millones", 7.5e8),
        ("$ 45 MM", 45e6),
        ("1,5 billones", 1.5e12),
        ("$12.000", 12000.0),
        ("12000", None),
        ("por estimar", None),
        ("", None),
        (4500000, 4.5e6),
    ],
)
def test_money_parser(text, amount):
    got = strategy_service.parse_money(text)
    if amount is None:
        assert got is None
    else:
        assert got == pytest.approx(amount, rel=1e-12)


# --------------------------------------------------------------------------- #
#  2. Agregados del tablero contra la verdad de referencia
# --------------------------------------------------------------------------- #
def test_every_aggregate_matches_reference_across_filter_combinations(db):
    today = ttm.today_co()
    cycle = _cycle(db, "Ciclo P-1")
    truth = seed_cycle(db, cycle, 160, seed=17, as_of=today)
    rows = strategy_service.collect_rows(db, cycle)
    assert len(rows) == 160

    cluster_opts = [None, 0] + [c.id for c in db.query(Cluster).all()][:3]
    type_opts = [None, 0] + [t.id for t in db.query(TechType).all()][:2]
    band_opts = ["", "inminente", "transicion", "desconocido", "lejano"]
    phase_opts = ["", "fase_iii", "autorizada", "sin_fase"]
    stage_opts = ["", "filtered", "prioritized", "published"]
    pmin_opts = [None, 0, 4, 6]
    date_opts = [(None, None), (dt.date(2026, 3, 1), None), (None, dt.date(2026, 5, 31)), (dt.date(2026, 2, 1), dt.date(2026, 2, 28))]
    checked = 0
    rng = random.Random(5)
    combos = list(itertools.product(cluster_opts, type_opts, band_opts, phase_opts, stage_opts, pmin_opts, date_opts))
    rng.shuffle(combos)
    for cluster_id, type_id, band, phase, stage, pmin, (dfrom, dto) in combos[:900]:
        f = dict(cluster_id=cluster_id, tech_type_id=type_id, band=band, phase=phase, stage=stage, priority_min=pmin, date_from=dfrom, date_to=dto)
        payload = strategy_service.build_payload(db, cycle, rows=rows, **f)
        assert_matches(payload, ref_aggregates(ref_filter(truth, f)), str(f))
        checked += 1
    # Filtro por estado, uno por uno.
    for status in ALL_STATUSES:
        payload = strategy_service.build_payload(db, cycle, rows=rows, status=status)
        assert_matches(payload, ref_aggregates(ref_filter(truth, {"status": status})), status)
        checked += 1
    assert checked == 900 + len(ALL_STATUSES)


def test_funnel_is_monotonic_and_ignores_status_from_other_cycles(db):
    """Bug corregido: el embudo contaba `Technology.status` global.

    Una tecnologia publicada en el ciclo IV aparecia como evaluada y publicada en
    el ciclo I, donde quedo bajo vigilancia: evaluadas > priorizadas.
    """
    c1 = _cycle(db, "Ciclo A", status="en_evaluacion")
    c4 = _cycle(db, "Ciclo B", status="en_evaluacion")
    t = Technology(commercial_name="Pelabresib", status="publicada")
    db.add(t)
    db.commit()
    db.add(CycleTechnology(cycle_id=c1.id, technology_id=t.id, status="bajo_vigilancia", priority_points=3))
    db.add(CycleTechnology(cycle_id=c4.id, technology_id=t.id, status="publicada", priority_points=5))
    db.commit()
    f1 = strategy_service.dashboard_for(db, c1, include_restricted=False)["funnel"]
    assert f1 == {"captured": 1, "filtered": 1, "prioritized": 0, "evaluated": 0, "published": 0}
    f4 = strategy_service.dashboard_for(db, c4, include_restricted=False)["funnel"]
    assert f4 == {"captured": 1, "filtered": 1, "prioritized": 1, "evaluated": 1, "published": 1}


def test_published_report_counts_even_if_entry_stayed_in_evaluation(db):
    cycle = _cycle(db, "Ciclo D", status="en_evaluacion")
    a = Technology(commercial_name="A")
    b = Technology(commercial_name="B")
    c = Technology(commercial_name="C")
    db.add_all([a, b, c])
    db.commit()
    db.add_all(
        [
            CycleTechnology(cycle_id=cycle.id, technology_id=a.id, status="en_evaluacion", priority_points=5),
            CycleTechnology(cycle_id=cycle.id, technology_id=b.id, status="priorizada", priority_points=4),
            CycleTechnology(cycle_id=cycle.id, technology_id=c.id, status="priorizada", priority_points=4),
            EvaluationDoc(cycle_id=cycle.id, technology_id=a.id, status="publicado", body={}),
            EvaluationDoc(cycle_id=cycle.id, technology_id=b.id, status="borrador", body={}),
        ]
    )
    db.commit()
    data = strategy_service.dashboard_for(db, cycle, include_restricted=False)
    assert data["funnel"] == {"captured": 3, "filtered": 3, "prioritized": 3, "evaluated": 2, "published": 1}
    assert data["conversion"] == {"filter_rate": 100, "priority_rate": 100, "eval_rate": 67, "publish_rate": 33}
    assert {r["technology_id"]: r["stage"] for r in data["ttm_scatter"]} == {
        a.id: "published",
        b.id: "evaluated",
        c.id: "prioritized",
    }


def test_conversion_rounds_half_up(db):
    cycle = _cycle(db, "Ciclo R")
    techs = [Technology(commercial_name=f"R{i}") for i in range(8)]
    db.add_all(techs)
    db.commit()
    for i, t in enumerate(techs):
        db.add(CycleTechnology(cycle_id=cycle.id, technology_id=t.id, status="filtrada_apta_priorizacion" if i < 5 else "asignada_a_ciclo"))
    db.commit()
    data = strategy_service.dashboard_for(db, cycle, include_restricted=False)
    assert data["conversion"]["filter_rate"] == 63  # 62,5 % sube a 63 (antes: 62)


def test_captured_on_uses_colombia_date(db):
    cycle = _cycle(db, "Ciclo TZ")
    t = Technology(commercial_name="Medianoche", captured_at=dt.datetime(2026, 3, 1, 3, 0, tzinfo=dt.timezone.utc))
    db.add(t)
    db.commit()
    db.add(CycleTechnology(cycle_id=cycle.id, technology_id=t.id, status="asignada_a_ciclo"))
    db.commit()
    row = strategy_service.collect_rows(db, cycle)[0]
    assert row["captured_on"] == "2026-02-28"
    feb = strategy_service.dashboard_for(db, cycle, include_restricted=False, date_from=dt.date(2026, 2, 28), date_to=dt.date(2026, 2, 28))
    assert feb["funnel"]["captured"] == 1
    mar = strategy_service.dashboard_for(db, cycle, include_restricted=False, date_from=dt.date(2026, 3, 1))
    assert mar["funnel"]["captured"] == 0


# --------------------------------------------------------------------------- #
#  3. Datamart al cierre
# --------------------------------------------------------------------------- #
def test_datamart_at_close_serves_filters_without_touching_transactions(db):
    today = ttm.today_co()
    cycle = _cycle(db, "Ciclo DM", status="en_evaluacion")
    truth = seed_cycle(db, cycle, 90, seed=23, as_of=today)
    # Sin bloqueos de cierre: toda tecnologia en evaluacion lleva justificacion.
    for entry in db.query(CycleTechnology).filter(CycleTechnology.cycle_id == cycle.id, CycleTechnology.status == "en_evaluacion"):
        entry.exclusion_note = "Cierre con justificacion de prueba"
    db.commit()
    cycle_service.close_cycle(db, cycle, actor="admin@iets.org.co")

    dm = db.query(CycleDatamart).filter(CycleDatamart.cycle_id == cycle.id).one()
    assert dm.payload["schema"] == strategy_service.DATAMART_SCHEMA
    assert len(dm.payload["rows"]) == 90
    assert db.query(TimeToMarketSnapshot).filter(TimeToMarketSnapshot.cycle_id == cycle.id).count() == 90

    for f in [{}, {"band": "inminente"}, {"stage": "prioritized", "priority_min": 4}, {"cluster_id": 0}, {"phase": "fase_iii", "date_from": dt.date(2026, 4, 1)}]:
        data = strategy_service.dashboard_for(db, cycle, include_restricted=True, **f)
        assert data["from_cache"] is True, f
        assert data["refreshed_at"] is not None
        assert_matches(data, ref_aggregates(ref_filter(truth, f)), str(f))

    # Cambiar datos transaccionales despues del cierre no altera el datamart.
    before = strategy_service.dashboard_for(db, cycle, include_restricted=False)
    moved = db.get(Technology, truth[0]["technology_id"])
    moved.cluster_id = None
    moved.fda_approval_date = today + dt.timedelta(days=5000)
    db.commit()
    after = strategy_service.dashboard_for(db, cycle, include_restricted=False)
    assert after["by_cluster"] == before["by_cluster"]
    assert after["by_band"] == before["by_band"]
    # Al refrescar, si.
    strategy_service.snapshot_ttm(db, cycle)
    strategy_service.refresh_datamart(db, cycle)
    db.commit()
    refreshed = strategy_service.dashboard_for(db, cycle, include_restricted=False)
    assert {r["technology_id"]: r["band"] for r in refreshed["ttm_scatter"]}[moved.id] == "lejano"
    assert {r["technology_id"]: r["cluster"] for r in refreshed["ttm_scatter"]}[moved.id] == "Sin clúster"


def test_closed_cycle_with_legacy_datamart_is_computed_live(db):
    cycle = _cycle(db, "Ciclo L", status="cerrado_consolidado")
    t = Technology(commercial_name="Legado")
    db.add(t)
    db.commit()
    db.add(CycleTechnology(cycle_id=cycle.id, technology_id=t.id, status="priorizada", priority_points=4))
    db.add(CycleDatamart(cycle_id=cycle.id, payload={"funnel": {"captured": 99}}))
    db.commit()
    data = strategy_service.dashboard_for(db, cycle, include_restricted=False)
    assert data["from_cache"] is False
    assert data["funnel"]["captured"] == 1


# --------------------------------------------------------------------------- #
#  4. API: capa restringida, validacion y exportacion
# --------------------------------------------------------------------------- #
def _small_cycle(h: Harness) -> Cycle:
    db = h.db
    cycle = _cycle(db, "Ciclo API", status="en_evaluacion")
    cluster = db.query(Cluster).first()
    a = Technology(commercial_name="=HYPERLINK(\"x\")", cluster_id=cluster.id, phase3_completion_date=ttm.today_co())
    b = Technology(commercial_name="Beta", cluster_id=cluster.id)
    db.add_all([a, b])
    db.commit()
    db.add_all(
        [
            CycleTechnology(cycle_id=cycle.id, technology_id=a.id, status="en_evaluacion", priority_points=6, priority_pct=100),
            CycleTechnology(cycle_id=cycle.id, technology_id=b.id, status="filtrada_apta_priorizacion"),
            EvaluationDoc(
                cycle_id=cycle.id,
                technology_id=a.id,
                status="borrador",
                body={"budget_year_1": "3.200 millones", "budget_year_2": "", "budget_year_3": "por estimar", "comparators_sgsss": "Imatinib"},
            ),
        ]
    )
    db.commit()
    return cycle


def test_api_public_layer_never_carries_budget(h):
    cycle = _small_cycle(h)
    for role in ("evaluador_tecnico", "evaluador_clinico", "revisor_pares"):
        client = h.as_role(role)
        body = client.get("/api/strategy/dashboard", params={"cycle_id": cycle.id}).json()
        assert body["restricted"] is False
        for key in ("budget_heatmap", "budget_items", "budget_unparsed", "comparators"):
            assert body.get(key) is None, (role, key)
        for row in body["ttm_scatter"]:
            assert not {"year1", "year2", "year3", "comparators", "budget_unparsed"} & set(row), role
        assert "3200000000" not in client.get("/api/strategy/dashboard", params={"cycle_id": cycle.id}).text
        graph = client.get("/api/strategy/graph", params={"cycle_id": cycle.id}).json()
        assert not [n for n in graph["nodes"] if n["type"] in ("budget", "comparators")], role
        csv_text = client.get("/api/strategy/dashboard/export", params={"cycle_id": cycle.id, "format": "csv"}).content.decode("utf-8-sig")
        assert "Impacto presupuestal" not in csv_text and "3200000000" not in csv_text

    for role in ("superadmin", "tomador_decisiones"):
        body = h.as_role(role).get("/api/strategy/dashboard", params={"cycle_id": cycle.id}).json()
        assert body["restricted"] is True
        assert body["budget_heatmap"][0]["year1"] == 3.2e9
        assert body["budget_unparsed"] == 1  # "por estimar" no es un monto
        assert body["comparators"][0]["comparators"] == "Imatinib"


def test_api_rejects_unknown_filter_values(h):
    cycle = _small_cycle(h)
    client = h.as_role("superadmin")
    for params in (
        {"band": "pronto"},
        {"phase": "fase_v"},
        {"status": "inventado"},
        {"stage": "casi"},
        {"date_from": "2026-05-01", "date_to": "2026-04-01"},
        {"cluster_id": -1},
        {"priority_min": -2},
    ):
        res = client.get("/api/strategy/dashboard", params={"cycle_id": cycle.id, **params})
        assert res.status_code == 422, params
    assert client.get("/api/strategy/dashboard", params={"cycle_id": 99999}).status_code == 404


def test_api_export_csv_and_xlsx_match_the_cut(h):
    from openpyxl import load_workbook

    cycle = _small_cycle(h)
    client = h.as_role("tomador_decisiones")
    res = client.get("/api/strategy/dashboard/export", params={"cycle_id": cycle.id, "format": "csv", "stage": "prioritized"})
    assert res.status_code == 200
    assert "attachment" in res.headers["content-disposition"] and ".csv" in res.headers["content-disposition"]
    rows = list(csv.reader(io.StringIO(res.content.decode("utf-8-sig")), delimiter=";"))
    header, data = rows[0], rows[1:]
    assert "Impacto presupuestal anio 1" in header
    assert len(data) == 1  # solo la priorizada
    record = dict(zip(header, data[0]))
    assert record["Tecnologia"].startswith("'=")  # formula neutralizada
    assert record["Puntos P1-P6"] == "6"
    assert record["Impacto presupuestal anio 1"] == "3200000000"
    assert record["Meses al mercado"] == f"{round(180 / 30.44, 2):.2f}".rstrip("0").rstrip(".").replace(".", ",")
    assert record["Etapa alcanzada del embudo"] == "Evaluadas"

    xres = client.get("/api/strategy/dashboard/export", params={"cycle_id": cycle.id, "format": "xlsx"})
    assert xres.status_code == 200
    wb = load_workbook(io.BytesIO(xres.content))
    sheet = wb["Recorte"]
    values = list(sheet.values)
    assert len(values) == 1 + 2
    xheader = list(values[0])
    points_col = xheader.index("Puntos P1-P6")
    by_name = {r[1]: r for r in values[1:]}
    assert by_name["Beta"][points_col] is None  # sin calificar, no 0
    assert wb["Contexto"]["B4"].value == 2

    tech = h.as_role("evaluador_tecnico").get("/api/strategy/dashboard/export", params={"cycle_id": cycle.id, "format": "xlsx"})
    theader = list(load_workbook(io.BytesIO(tech.content))["Recorte"].values)[0]
    assert not [c for c in theader if "presupuestal" in str(c)]
    assert h.client.get("/api/strategy/dashboard/export", params={"cycle_id": cycle.id, "format": "pdf"}).status_code == 422


# --------------------------------------------------------------------------- #
#  5. Bandeja por perfil (GET /api/dashboard/my-work)
# --------------------------------------------------------------------------- #
EXPECTED_KEYS = {
    "superadmin": ["rate", "classify", "novelty", "filter_ready", "invima", "drafts", "submissions", "merges", "cycle_deadline", "close_blockers", "bulletins_pending"],
    "evaluador_tecnico": ["rate", "classify", "novelty", "filter_ready", "invima"],
    "evaluador_clinico": ["rate", "drafts"],
    "tomador_decisiones": ["bulletins_published", "alerts", "dashboard"],
    "revisor_pares": ["reviews"],
}


def _seed_work(h: Harness) -> dict:
    db = h.db
    today = ttm.today_co()
    now = dt.datetime.now(dt.timezone.utc)
    cluster = db.query(Cluster).first()
    ttype = db.query(TechType).first()
    cycle = _cycle(db, "Ciclo W", data_cutoff_on=today + dt.timedelta(days=5), bulletin_due_on=today + dt.timedelta(days=20))

    def tech(name, **kw):
        # Solo las senales de staging quedan "capturada_no_asignada".
        kw.setdefault("status", "asignada_a_ciclo")
        t = Technology(commercial_name=name, **kw)
        db.add(t)
        db.flush()
        return t

    def entry(t, status, **kw):
        e = CycleTechnology(cycle_id=cycle.id, technology_id=t.id, status=status, **kw)
        db.add(e)
        return e

    a, b, c, d = tech("A apta sin calificar"), tech("B apta tecnico listo"), tech("C congelada"), tech("D completa")
    entry(a, "filtrada_apta_priorizacion")
    entry(b, "filtrada_apta_priorizacion")
    entry(c, "bajo_vigilancia", frozen=True)
    entry(d, "priorizada", priority_points=6, priority_pct=100)
    for code in ("P1", "P5", "P6"):
        db.add(PriorityScore(cycle_id=cycle.id, technology_id=b.id, criterion=code, value=1))
    for code in ("P1", "P2", "P3", "P4", "P5", "P6"):
        db.add(PriorityScore(cycle_id=cycle.id, technology_id=d.id, criterion=code, value=1))

    e_, f_, g_ = tech("E sin novedad"), tech("F sin INVIMA"), tech("G lista")
    entry(e_, "asignada_a_ciclo")
    entry(f_, "asignada_a_ciclo")
    entry(g_, "asignada_a_ciclo")
    db.add(NoveltyAssessment(cycle_id=cycle.id, technology_id=f_.id, option_code="no_disponible_en_pais"))
    db.add(NoveltyAssessment(cycle_id=cycle.id, technology_id=g_.id, option_code="no_disponible_en_pais", invima_checked_at=now, has_valid_registry=False))

    hx, ix = tech("H en evaluacion"), tech("I con observaciones")
    entry(hx, "en_evaluacion", priority_points=5)
    entry(ix, "en_evaluacion", priority_points=5)
    doc_h = EvaluationDoc(cycle_id=cycle.id, technology_id=hx.id, status="borrador", title="Ficha H", body={})
    doc_i = EvaluationDoc(cycle_id=cycle.id, technology_id=ix.id, status="con_observaciones", title="Informe I", body={})
    db.add_all([doc_h, doc_i])

    closed = _cycle(db, "Ciclo cerrado W", status="cerrado_consolidado", opened_on=today - dt.timedelta(weeks=30))
    closed.closed_at = now - dt.timedelta(days=10)
    old = tech("Vieja")
    db.add(EvaluationDoc(cycle_id=closed.id, technology_id=old.id, status="borrador", body={}))
    for name, status in (("K1", "priorizada"), ("K2", "publicada"), ("K3", "no_priorizada")):
        db.add(CycleTechnology(cycle_id=closed.id, technology_id=tech(name).id, status=status))

    tech("S1 sin cluster", status="capturada_no_asignada", tech_type_id=ttype.id)
    tech("S2 sin tipologia", status="capturada_no_asignada", cluster_id=cluster.id)
    tech("S3 completa", status="capturada_no_asignada", cluster_id=cluster.id, tech_type_id=ttype.id)
    tech("S4 fusionada", status="capturada_no_asignada", merged_into_id=a.id)

    db.add_all(
        [
            Submission(commercial_name="Post 1", status="recibida"),
            Submission(commercial_name="Post 2", status="recibida"),
            Submission(commercial_name="Post 3", status="aceptada"),
        ]
    )
    db.flush()
    db.add(MergeProposal(technology_a_id=min(a.id, b.id), technology_b_id=max(a.id, b.id), status="propuesta", score=91))
    db.add(MergeProposal(technology_a_id=min(c.id, d.id), technology_b_id=max(c.id, d.id), status="confirmada", score=95))
    db.add_all(
        [
            Bulletin(cycle_id=closed.id, title="Boletin pendiente", status="pendiente_aprobacion"),
            Bulletin(cycle_id=closed.id, title="Boletin reciente", status="publicado", published_at=now - dt.timedelta(days=3)),
            Bulletin(cycle_id=closed.id, title="Boletin viejo", status="publicado", published_at=now - dt.timedelta(days=200)),
        ]
    )
    tomador = h.users["tomador_decisiones"]
    db.add_all(
        [
            AlertEvent(user_id=tomador.id, kind="k", title="Alerta 1"),
            AlertEvent(user_id=tomador.id, kind="k", title="Alerta 2"),
            AlertEvent(user_id=tomador.id, kind="k", title="Leida", read_at=now),
            AlertEvent(user_id=h.users["superadmin"].id, kind="k", title="De otro"),
        ]
    )
    db.flush()
    revisor = h.users["revisor_pares"]
    db.add_all(
        [
            ReviewAssignment(doc_id=doc_h.id, kind="interno", reviewer_user_id=revisor.id, reviewer_email=revisor.email, status="invitado"),
            ReviewAssignment(doc_id=doc_i.id, kind="externo", reviewer_email=revisor.email.upper(), status="en_lectura"),
            ReviewAssignment(doc_id=doc_i.id, kind="interno", reviewer_user_id=revisor.id, status="aprobado"),
            ReviewAssignment(doc_id=doc_h.id, kind="externo", reviewer_email=revisor.email, status="revocado"),
        ]
    )
    source = Source(title="Fuente W", url="https://w.example")
    db.add(source)
    db.flush()

    def finding(title, score, status="nuevo"):
        f = Finding(source_id=source.id, title=title, content_hash=title, screening_score=score, status=status)
        db.add(f)
        db.flush()
        return f

    f1 = finding("Senal sin tecnologia", 90)
    f2 = finding("Senal en staging", 80)
    f3 = finding("Senal ya publicada", 95)
    f4 = finding("Senal fusionada", 99)
    finding("Senal descartada", 100, status="descartado")
    f6 = finding("Senal asignada", 70)
    t2 = tech("T2", status="capturada_no_asignada", finding_id=f2.id, cluster_id=cluster.id, tech_type_id=ttype.id)
    tech("T3", status="publicada", finding_id=f3.id)
    tech("T4", status="capturada_no_asignada", finding_id=f4.id, merged_into_id=a.id, cluster_id=cluster.id, tech_type_id=ttype.id)
    t6 = tech("T6", status="asignada_a_ciclo", finding_id=f6.id)
    db.commit()
    return {"cycle": cycle, "f1": f1, "t2": t2, "t6": t6, "g": g_, "e": e_, "f": f_}


def test_my_work_counts_per_profile(h):
    ctx = _seed_work(h)
    expected_counts = {
        "superadmin": {"rate": 2, "classify": 2, "novelty": 2, "filter_ready": 1, "invima": 0, "drafts": 2, "submissions": 2, "merges": 1, "cycle_deadline": 5, "close_blockers": 2, "bulletins_pending": 1},
        "evaluador_tecnico": {"rate": 1, "classify": 2, "novelty": 2, "filter_ready": 1, "invima": 0},
        "evaluador_clinico": {"rate": 2, "drafts": 2},
        "tomador_decisiones": {"bulletins_published": 1, "alerts": 2, "dashboard": 2},
        "revisor_pares": {"reviews": 2},
    }
    for role, keys in EXPECTED_KEYS.items():
        body = h.as_role(role).get("/api/dashboard/my-work").json()
        assert body["role"] == role
        got = {i["key"]: i for i in body["items"]}
        assert [i["key"] for i in body["items"]] == keys, role
        assert {k: v["count"] for k, v in got.items()} == expected_counts[role], role
        for item in body["items"]:
            assert item["to"].startswith("/") and item["help"] and item["action_label"], (role, item["key"])
            assert len(item["samples"]) == min(item["count"], 5) or item["key"] in ("invima", "cycle_deadline", "dashboard"), (role, item["key"])
        actionable = sum(i["count"] for i in body["items"] if i["severity"] == "accion")
        assert body["total_pending"] == actionable
        # La bandeja y el KPI "Me toca calificar" usan la misma regla.
        if "rate" in got:
            wb = h.client.get("/api/dashboard/workbench").json()
            assert wb["cycle_pending_for_me"] == got["rate"]["count"], role

    admin = {i["key"]: i for i in h.as_role("superadmin").get("/api/dashboard/my-work").json()["items"]}
    assert admin["rate"]["label"].startswith("P1, P2, P3, P4, P5, P6")
    tecnico = {i["key"]: i for i in h.as_role("evaluador_tecnico").get("/api/dashboard/my-work").json()["items"]}
    assert tecnico["rate"]["label"] == "P1, P5, P6 por calificar"
    assert tecnico["rate"]["samples"][0]["detail"] == "Falta P1, P5, P6"
    details = {s["id"]: s["detail"] for s in tecnico["novelty"]["samples"]}
    assert details == {ctx["e"].id: "Falta vía de novedad", ctx["f"].id: "Falta cruce con INVIMA"}
    assert [s["id"] for s in tecnico["filter_ready"]["samples"]] == [ctx["g"].id]
    assert tecnico["invima"]["severity"] == "aviso" and tecnico["invima"]["label"] == "Índice INVIMA vacío"
    assert admin["cycle_deadline"]["label"].endswith("corte de datos en 5 día(s)")
    clinico = {i["key"]: i for i in h.as_role("evaluador_clinico").get("/api/dashboard/my-work").json()["items"]}
    assert clinico["rate"]["label"] == "P2, P3, P4 por calificar"


def test_work_queue_skips_finished_signals_and_links_to_the_action(h):
    ctx = _seed_work(h)
    body = h.as_role("evaluador_tecnico").get("/api/dashboard/my-work").json()
    assert body["queue_total"] == 3
    assert [q["title"] for q in body["queue"]] == ["Senal sin tecnologia", "Senal en staging", "Senal asignada"]
    assert [q["to"] for q in body["queue"]] == [
        "/senales",
        f"/bandeja-entrada?tecnologia={ctx['t2'].id}",
        f"/filtrado?tecnologia={ctx['t6'].id}",
    ]
    assert h.as_role("tomador_decisiones").get("/api/dashboard/my-work").json()["queue"] == []


def test_deadline_overdue_and_far(h):
    db = h.db
    today = ttm.today_co()
    cycle = _cycle(db, "Ciclo V", data_cutoff_on=today - dt.timedelta(days=9), bulletin_due_on=today - dt.timedelta(days=3))
    items = {i["key"]: i for i in h.as_role("superadmin").get("/api/dashboard/my-work").json()["items"]}
    assert items["cycle_deadline"]["label"].endswith("boletín vencido hace 3 día(s)")
    cycle.data_cutoff_on = today + dt.timedelta(days=30)
    cycle.bulletin_due_on = today + dt.timedelta(days=40)
    db.commit()
    items = {i["key"]: i for i in h.as_role("superadmin").get("/api/dashboard/my-work").json()["items"]}
    assert "cycle_deadline" not in items


# --------------------------------------------------------------------------- #
#  6. Rendimiento: p95 < 2 s con el volumen de tres ciclos (3 x 300)
# --------------------------------------------------------------------------- #
def test_dashboard_p95_under_two_seconds_with_three_cycles(h):
    db = h.db
    today = ttm.today_co()
    cycles = []
    for k, status in enumerate(("en_evaluacion", "en_priorizacion", "en_evaluacion")):
        cycle = _cycle(db, f"Ciclo PERF-{k}", status=status)
        seed_cycle(db, cycle, 300, seed=100 + k, as_of=today)
        cycles.append(cycle)
    first = cycles[0]
    for entry in db.query(CycleTechnology).filter(CycleTechnology.cycle_id == first.id, CycleTechnology.status == "en_evaluacion"):
        entry.exclusion_note = "Cierre de prueba"
    db.commit()
    cycle_service.close_cycle(db, first, actor="admin@iets.org.co")

    client = h.as_role("superadmin")
    scenarios = [{}, {"band": "inminente"}, {"stage": "prioritized", "priority_min": 4}, {"cluster_id": 1, "phase": "fase_iii"}]
    timings = []
    for _ in range(5):
        for cycle in cycles:
            for params in scenarios:
                start = time.perf_counter()
                res = client.get("/api/strategy/dashboard", params={"cycle_id": cycle.id, **params})
                timings.append(time.perf_counter() - start)
                assert res.status_code == 200
                assert res.json()["total_in_cycle"] == 300
    timings.sort()
    p95 = timings[int(len(timings) * 0.95) - 1]
    print(f"\n[frente C] tablero 3x300: n={len(timings)} p50={statistics.median(timings) * 1000:.1f} ms p95={p95 * 1000:.1f} ms max={timings[-1] * 1000:.1f} ms")
    assert p95 < 2.0


# --------------------------------------------------------------------------- #
#  7. Capa publica (/transparencia): mismo criterio que la lista publica
# --------------------------------------------------------------------------- #
def test_public_stats_use_the_same_criterion_as_the_public_list(db):
    """Bug corregido (reportado por el frente B): /transparencia contaba
    `Technology.status == "publicada"`, incluidas las confidenciales, mientras
    que la lista publica las excluye. El conteo revelaba que existian."""
    c1 = _cycle(db, "Ciclo PUB-1", status="cerrado_consolidado")
    c2 = _cycle(db, "Ciclo PUB-2", status="cerrado_consolidado")
    clusters = db.query(Cluster).order_by(Cluster.id).all()
    now = dt.datetime.now(dt.timezone.utc)
    a = Technology(commercial_name="Publica A", cluster_id=clusters[0].id, status="publicada")
    b = Technology(commercial_name="Confidencial B", cluster_id=clusters[0].id, status="publicada")
    c = Technology(commercial_name="Dos ciclos C", cluster_id=clusters[1].id, status="publicada")
    d = Technology(commercial_name="Sin cluster D", status="publicada")
    e = Technology(commercial_name="Estado sin expediente E", cluster_id=clusters[1].id, status="publicada")
    f = Technology(commercial_name="Borrador F", cluster_id=clusters[1].id, status="en_evaluacion")
    db.add_all([a, b, c, d, e, f])
    db.commit()
    db.add_all(
        [
            EvaluationDoc(cycle_id=c1.id, technology_id=a.id, status="publicado", confidential=False, published_at=now, body={}),
            EvaluationDoc(cycle_id=c1.id, technology_id=b.id, status="publicado", confidential=True, published_at=now, body={}),
            EvaluationDoc(cycle_id=c1.id, technology_id=c.id, status="publicado", confidential=False, published_at=now, body={}),
            EvaluationDoc(cycle_id=c2.id, technology_id=c.id, status="publicado", confidential=False, published_at=now, body={}),
            EvaluationDoc(cycle_id=c1.id, technology_id=d.id, status="publicado", confidential=False, published_at=now, body={}),
            EvaluationDoc(cycle_id=c1.id, technology_id=f.id, status="borrador", confidential=False, body={}),
        ]
    )
    db.commit()
    listed = strategy_service.list_public_fiches(db, limit=50)
    stats = strategy_service.public_stats(db)
    assert sorted(i["id"] for i in listed) == sorted([a.id, c.id, d.id])
    assert stats["published"] == len(listed) == 3
    assert sum(i["value"] for i in stats["by_cluster"]) == stats["published"]
    assert {i["label"]: i["value"] for i in stats["by_cluster"]} == {
        clusters[0].name: 1,
        clusters[1].name: 1,
        "Sin clúster": 1,
    }
    # Por cluster, la lista publica y la estadistica coinciden.
    for cluster in clusters[:2]:
        in_list = len(strategy_service.list_public_fiches(db, cluster_id=cluster.id, limit=50))
        in_stats = {i["label"]: i["value"] for i in stats["by_cluster"]}.get(cluster.name, 0)
        assert in_list == in_stats
    assert "Confidencial B" not in str(stats)
